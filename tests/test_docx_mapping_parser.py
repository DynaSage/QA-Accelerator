from pathlib import Path

from scripts.create_samples import write_stm_docx_sample
from src.parsers.excel_mapping_parser import parse_mapping_workbook


def test_parse_stm_docx(tmp_path: Path):
    path = tmp_path / "stm.docx"
    write_stm_docx_sample(path)
    payload = parse_mapping_workbook(path)

    assert payload["target_table"] == "Person_Account"
    assert len(payload["mapping_rows"]) == 4

    first = payload["mapping_rows"][0]
    assert first["target_table"] == "Person_Account"
    assert first["target_column"] == "account_number"
    assert first["source_column"] == "account_number"
    assert "hldu_hldr_j" in first["source_table"]
    assert first["transformation"] == "DIRECT"

    concat = payload["mapping_rows"][1]
    assert concat["target_column"] == "combined_key"
    assert "CONCAT_WS" in concat["transformation"]

    lookup = payload["mapping_rows"][2]
    assert lookup["target_table"] == "Positions"
    assert lookup["source_column"] == "financial_account_id"
    assert "Lookup" in lookup["transformation"]

    assert any("Join person accounts" in item.get("pseudo_code", "") for item in payload["entity_logic"])


def test_parse_real_abc1234_stm_docx_if_present():
    from src.parsers.stm_extract import parse_stm_docx

    path = Path(r"C:\Users\utsanjan.maity\Downloads\ABC1234 Data Warehouse Documentation_V4 8.docx")
    if not path.exists():
        return
    payload = parse_stm_docx(path, write_excel=False)
    tables = {row["target_table"] for row in payload["mapping_rows"]}
    assert len(payload["mapping_rows"]) >= 200
    assert payload["document_kind"] == "word_stm"
    assert len(payload["stm_entities"]) == 17
    assert "Person_Account" in tables
    assert "Financial_Account" in tables
    assert "Positions" in tables
    assert payload["entity_logic"]
    assert "Environments" not in tables
    assert "Orchestrators" not in tables
    assert "Document_Information" not in tables


def test_parse_stm_docx_does_not_use_openpyxl(tmp_path: Path, monkeypatch):
    path = tmp_path / "stm.docx"
    write_stm_docx_sample(path)

    def _fail_excel(*args, **kwargs):
        raise AssertionError("openpyxl should not open a Word mapping document")

    monkeypatch.setattr("src.parsers.excel_mapping_parser.load_workbook", _fail_excel)
    from src.parsers.mapping_file import parse_mapping_file

    payload = parse_mapping_file(path)
    assert payload["mapping_rows"]


def test_word_stm_misnamed_as_xlsx_uses_docx_parser(tmp_path: Path, monkeypatch):
    source = tmp_path / "stm.docx"
    write_stm_docx_sample(source)
    misnamed = tmp_path / "mapping.xlsx"
    misnamed.write_bytes(source.read_bytes())

    def _fail_excel(*args, **kwargs):
        raise AssertionError("openpyxl should not open a Word document named .xlsx")

    monkeypatch.setattr("src.parsers.excel_mapping_parser.load_workbook", _fail_excel)
    payload = parse_mapping_workbook(misnamed)
    assert payload["mapping_rows"]


def test_stm_source_column_and_logic_header(tmp_path: Path):
    from docx import Document
    from src.parsers.stm_extract import canonical_stm_header, parse_stm_docx

    headers = ["Target Column", "Source Column & Logic", "Source Table(s)", "Original Source"]
    assert canonical_stm_header(headers[0]) == "target_column"
    assert canonical_stm_header(headers[1]) == "source_column"
    assert canonical_stm_header(headers[2]) == "source_table"
    assert canonical_stm_header(headers[3]) == "notes"

    path = tmp_path / "stm_logic.docx"
    document = Document()
    document.add_heading("Person Account:", 2)
    table = document.add_table(rows=2, cols=4)
    table.rows[0].cells[0].text = "Target Column"
    table.rows[0].cells[1].text = "Source Column & Logic"
    table.rows[0].cells[2].text = "Source Table(s)"
    table.rows[0].cells[3].text = "Original Source"
    table.rows[1].cells[0].text = "account_number"
    table.rows[1].cells[1].text = "account_number"
    table.rows[1].cells[2].text = "hldu_hldr_j"
    table.rows[1].cells[3].text = "SQL Server"
    document.save(path)

    payload = parse_stm_docx(path, write_excel=True)
    row = payload["mapping_rows"][0]
    assert row["target_column"] == "account_number"
    assert row["source_column"] == "account_number"
    assert row["source_table"] == "hldu_hldr_j"
    assert row["notes"] == "SQL Server"
    excel_path = Path(payload["normalized_excel_path"])
    assert excel_path.exists()
    from src.parsers.excel_mapping_parser import parse_mapping_workbook

    excel_payload = parse_mapping_workbook(excel_path)
    assert excel_payload["mapping_rows"][0]["source_column"] == "account_number"


def test_canonical_header_tokens_do_not_treat_sources_as_source_column():
    from src.parsers.stm_extract import canonical_stm_header, is_stm_table

    assert canonical_stm_header("Sources") is None
    assert canonical_stm_header("Source Column & Logic") == "source_column"
    assert canonical_stm_header("Source Column / Logic") == "source_column"
    assert is_stm_table(["Orchestrator", "Sources", "Schedule"]) is False
    assert is_stm_table(["Document ID", "Title", "Version"]) is False


def test_parse_stm_docx_skips_architecture_tables(tmp_path: Path, monkeypatch):
    from docx import Document
    from src.parsers.stm_extract import parse_stm_docx

    def _fail_llm(*args, **kwargs):
        raise AssertionError("architecture tables should not call the header LLM")

    monkeypatch.setattr("src.parsers.stm_extract._llm_header_map", _fail_llm)

    path = tmp_path / "warehouse_stm.docx"
    document = Document()
    document.add_heading("Document Information", 1)
    info = document.add_table(rows=2, cols=3)
    info.rows[0].cells[0].text = "Document ID"
    info.rows[0].cells[1].text = "Title"
    info.rows[0].cells[2].text = "Version"
    info.rows[1].cells[0].text = "ABC1234"
    info.rows[1].cells[1].text = "Data Warehouse Documentation"
    info.rows[1].cells[2].text = "4.8"

    document.add_heading("Source to Target Mapping (STM) & Pseudo Code", 1)
    document.add_heading("Person Account:", 2)
    document.add_heading("STM", 3)
    person = document.add_table(rows=2, cols=4)
    person.rows[0].cells[0].text = "Target Column"
    person.rows[0].cells[1].text = "Source Column & Logic"
    person.rows[0].cells[2].text = "Source Table(s)"
    person.rows[0].cells[3].text = "Original Source"
    person.rows[1].cells[0].text = "account_number"
    person.rows[1].cells[1].text = "account_number"
    person.rows[1].cells[2].text = "hldu_hldr_j"
    person.rows[1].cells[3].text = "SQL Server"

    document.add_heading("Environments", 2)
    env = document.add_table(rows=2, cols=3)
    env.rows[0].cells[0].text = "Environment"
    env.rows[0].cells[1].text = "Catalog"
    env.rows[0].cells[2].text = "Purpose"
    env.rows[1].cells[0].text = "DEV"
    env.rows[1].cells[1].text = "abc_dev"
    env.rows[1].cells[2].text = "Development"

    document.add_heading("Orchestrators", 2)
    orch = document.add_table(rows=2, cols=3)
    orch.rows[0].cells[0].text = "Orchestrator"
    orch.rows[0].cells[1].text = "Sources"
    orch.rows[0].cells[2].text = "Schedule"
    orch.rows[1].cells[0].text = "ADF"
    orch.rows[1].cells[1].text = "SQL Server"
    orch.rows[1].cells[2].text = "Daily"

    document.add_heading("Positions", 2)
    document.add_heading("Pseudo Code", 3)
    document.add_paragraph("Lookup financial_account_id from gold financial_account.")
    document.add_heading("STM", 3)
    positions = document.add_table(rows=2, cols=3)
    positions.rows[0].cells[0].text = "Target Column"
    positions.rows[0].cells[1].text = "Source Column / Logic"
    positions.rows[0].cells[2].text = "Source Table(s)"
    positions.rows[1].cells[0].text = "financial_account_id"
    positions.rows[1].cells[1].text = "financial_account_id (Lookup based on Account no)"
    positions.rows[1].cells[2].text = "financial_account"

    document.save(path)
    payload = parse_stm_docx(path, write_excel=False)

    tables = {row["target_table"] for row in payload["mapping_rows"]}
    assert payload["document_kind"] == "word_stm"
    assert tables == {"Person_Account", "Positions"}
    assert len(payload["mapping_rows"]) == 2
    assert {item["target_table"] for item in payload["stm_entities"]} == tables
    assert any("Lookup financial_account_id" in item.get("pseudo_code", "") for item in payload["entity_logic"])
    assert not any(row["source_column"] == "SQL Server" for row in payload["mapping_rows"])
    assert "Environments" not in tables
    assert "Orchestrators" not in tables
    assert "Document_Information" not in tables


def test_legacy_resolve_stm_headers_accepts_source_column_and_logic():
    from src.parsers.docx_mapping_parser import _resolve_stm_headers

    headers = ["Target Column", "Source Column & Logic", "Source Table(s)", "Original Source"]
    resolved = _resolve_stm_headers(headers)
    assert resolved["target_column"] == 0
    assert resolved["source_column"] == 1
    assert resolved["source_table"] == 2
