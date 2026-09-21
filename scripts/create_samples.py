from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font


ROOT = Path(__file__).resolve().parents[1]
SAMPLES = ROOT / "samples"
REQUIREMENTS = SAMPLES / "requirements"
MAPPINGS = SAMPLES / "mappings"
TEMPLATES = SAMPLES / "templates"


def write_requirement_samples() -> None:
    REQUIREMENTS.mkdir(parents=True, exist_ok=True)
    brd_text = """ETL Requirement Document: Customer Dimension Load

Objective:
Load customer data from staging into the warehouse dimension table for QA validation.

Source Table:
main.stg.customer_raw

Target Table:
main.dw.dim_customer

Primary Key:
customer_id

Load Type:
Incremental load using updated_at compared against the last successful watermark in main.ctl.etl_watermark.

Business Rules:
1. email_address must contain '@'
2. customer_id must never be NULL
3. is_active must be 0 or 1 only

Transformation Rules:
1. customer_name must be upper-cased and trimmed
2. email_address must be lower-cased
3. is_active is 1 when source status_cd = 'A', else 0
4. Soft-deleted source rows (status_cd = 'D') must not land in target

Not Null Columns in Target:
customer_id, customer_name, effective_from

Additional Notes:
All tables live in Azure Databricks Unity Catalog.
"""
    (REQUIREMENTS / "sample_brd.txt").write_text(brd_text, encoding="utf-8")

    try:
        from docx import Document

        document = Document()
        for line in brd_text.splitlines():
            document.add_paragraph(line)
        document.save(REQUIREMENTS / "sample_brd.docx")
    except Exception:
        pass


def write_mapping_workbook(path: Path, *, template: bool = False) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Mapping"
    headers = [
        "Source Database",
        "Source Schema",
        "Source Table",
        "Source Column",
        "Source Data Type",
        "Transformation",
        "Target Database",
        "Target Schema",
        "Target Table",
        "Target Column",
        "Target Data Type",
        "Business Rule",
        "Nullable",
        "Primary Key",
        "Notes",
    ]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    if not template:
        rows = [
            [
                "main", "stg", "customer_raw", "cust_id", "INT",
                "CAST(cust_id AS BIGINT)",
                "main", "dw", "dim_customer", "customer_id", "BIGINT",
                "Must be unique", "N", "Y", "Business key",
            ],
            [
                "", "", "", "cust_name", "VARCHAR(100)",
                "TRIM(UPPER(cust_name))",
                "", "", "", "customer_name", "VARCHAR(100)",
                "Mandatory", "N", "N", "",
            ],
            [
                "", "", "", "email", "VARCHAR(150)",
                "LOWER(TRIM(email))",
                "", "", "", "email_address", "VARCHAR(150)",
                "Must contain @", "Y", "N", "",
            ],
            [
                "", "", "", "status_cd", "CHAR(1)",
                "CASE WHEN status_cd = 'A' THEN 1 ELSE 0 END",
                "", "", "", "is_active", "INT",
                "1 when A else 0", "N", "N", "",
            ],
            [
                "", "", "", "created_dt", "VARCHAR(20)",
                "CAST(created_dt AS DATE)",
                "", "", "", "effective_from", "DATE",
                "Effective date", "N", "N", "",
            ],
        ]
        for row in rows:
            sheet.append(row)

    for column_cells in sheet.columns:
        width = min(max(len(str(cell.value or "")) for cell in column_cells) + 2, 28)
        sheet.column_dimensions[column_cells[0].column_letter].width = width

    workbook.save(path)


def write_stm_docx_sample(path: Path) -> None:
    from docx import Document

    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_heading("Source to Target Mapping (STM) & Pseudo Code", 1)
    document.add_heading("Person Account:", 2)
    document.add_heading("STM", 3)
    table = document.add_table(rows=3, cols=4)
    table.rows[0].cells[0].text = "Target Column"
    table.rows[0].cells[1].text = "Source Column"
    table.rows[0].cells[2].text = "Source Table"
    table.rows[0].cells[3].text = "Original Source"
    table.rows[1].cells[0].text = "account_number"
    table.rows[1].cells[1].text = "account_number"
    table.rows[1].cells[2].text = "hldu_hldr_j, account_combined_details"
    table.rows[1].cells[3].text = "SQL Server (Pershing)"
    table.rows[2].cells[0].text = "combined_key"
    table.rows[2].cells[1].text = "CONCAT_WS(' ', account_number, ibd_number)"
    table.rows[2].cells[2].text = "hldu_hldr_j"
    table.rows[2].cells[3].text = "SQL Server (Pershing)"
    document.add_heading("Pseudo Code", 3)
    document.add_paragraph("Join person accounts on account_number and ibd_number.")

    document.add_heading("Positions", 2)
    document.add_heading("STM", 3)
    positions = document.add_table(rows=3, cols=3)
    positions.rows[0].cells[0].text = "Target Column"
    positions.rows[0].cells[1].text = "Source Column / Logic"
    positions.rows[0].cells[2].text = "Source Table(s)"
    positions.rows[1].cells[0].text = "financial_account_id"
    positions.rows[1].cells[1].text = "financial_account_id (Lookup based on Account no)"
    positions.rows[1].cells[2].text = "financial_account (gold)"
    positions.rows[2].cells[0].text = "id"
    positions.rows[2].cells[1].text = "Incremental Logic: MAX(existing_id) + Row_Number"
    positions.rows[2].cells[2].text = "System Generated"
    document.save(path)


def main() -> None:
    write_requirement_samples()
    write_mapping_workbook(MAPPINGS / "sample_mapping.xlsx")
    write_mapping_workbook(TEMPLATES / "mapping_template.xlsx", template=True)
    write_stm_docx_sample(MAPPINGS / "sample_stm_mapping.docx")
    print(f"Sample files created under {SAMPLES}")


if __name__ == "__main__":
    main()
