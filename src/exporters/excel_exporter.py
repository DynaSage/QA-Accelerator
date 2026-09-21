from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font

from src.models.test_case import StructuredTestCase


def _save_workbook(workbook: Workbook, output_path: Path) -> Path:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path


def _write_sheet(sheet, headers: list[str], rows: list[list[Any]]) -> None:
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append(row)
    for column_cells in sheet.columns:
        max_length = max(len(str(cell.value or "")) for cell in column_cells)
        sheet.column_dimensions[column_cells[0].column_letter].width = min(max_length + 2, 80)


def export_test_cases_to_excel(
    test_cases: list[dict[str, Any]],
    output_path: Path,
    *,
    artifacts: dict[str, Any] | None = None,
) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "QA Test Cases"

    rows = [StructuredTestCase.model_validate(case).to_row() for case in test_cases]
    if rows:
        headers = list(rows[0].keys())
        _write_sheet(sheet, headers, [[row.get(header, "") for header in headers] for row in rows])
    else:
        _write_sheet(
            sheet,
            ["Test Case ID", "Test Scenario", "Priority", "Expected Result"],
            [],
        )

    artifacts = artifacts or {}
    if artifacts.get("test_scenario_summary"):
        summary_sheet = workbook.create_sheet("Scenario Summary")
        summary_rows = artifacts["test_scenario_summary"]
        headers = ["Scenario ID", "Scenario Name", "Test Type", "Priority", "Description"]
        _write_sheet(
            summary_sheet,
            headers,
            [
                [
                    row.get("scenario_id", ""),
                    row.get("scenario_name", ""),
                    row.get("test_type", ""),
                    row.get("priority", ""),
                    row.get("description", ""),
                ]
                for row in summary_rows
            ],
        )

    if artifacts.get("requirement_traceability_matrix"):
        rtm_sheet = workbook.create_sheet("RTM")
        rtm_rows = artifacts["requirement_traceability_matrix"]
        headers = ["Requirement ID", "Requirement Text", "Test Case IDs", "Coverage Status"]
        _write_sheet(
            rtm_sheet,
            headers,
            [
                [
                    row.get("requirement_id", ""),
                    row.get("requirement_text", ""),
                    ", ".join(row.get("test_case_ids") or []),
                    row.get("coverage_status", ""),
                ]
                for row in rtm_rows
            ],
        )

    if artifacts.get("test_coverage_matrix"):
        coverage_sheet = workbook.create_sheet("Coverage Matrix")
        coverage_rows = artifacts["test_coverage_matrix"]
        headers = ["Test Type", "Total Cases", "High Priority Cases", "Coverage Notes"]
        _write_sheet(
            coverage_sheet,
            headers,
            [
                [
                    row.get("test_type", ""),
                    row.get("total_cases", ""),
                    row.get("high_priority_cases", ""),
                    row.get("coverage_notes", ""),
                ]
                for row in coverage_rows
            ],
        )

    if artifacts.get("automation_recommendations"):
        automation_sheet = workbook.create_sheet("Automation")
        automation_rows = artifacts["automation_recommendations"]
        headers = ["Test Case ID", "Automation Candidate", "Reason"]
        _write_sheet(
            automation_sheet,
            headers,
            [
                [
                    row.get("test_case_id", ""),
                    row.get("automation_candidate", ""),
                    row.get("reason", ""),
                ]
                for row in automation_rows
            ],
        )

    if artifacts.get("mapping_analysis"):
        analysis = artifacts["mapping_analysis"]
        mapping_sheet = workbook.create_sheet("Mapping Analysis")
        headers = [
            "Mapping ID",
            "Source",
            "Target",
            "Transformation Type",
            "Logic",
            "Risk",
            "Gaps",
            "Rule IDs",
            "Requires Review",
        ]
        rows = []
        for item in analysis.get("mappings") or []:
            source = item.get("source") or {}
            target = item.get("target") or {}
            transform = item.get("transformation") or {}
            rules = item.get("validation_rules") or []
            rows.append(
                [
                    item.get("mapping_id", ""),
                    f"{source.get('table', '')}.{source.get('column', '')}",
                    f"{target.get('table', '')}.{target.get('column', '')}",
                    transform.get("type", ""),
                    transform.get("logic", ""),
                    item.get("risk", ""),
                    "; ".join(item.get("gaps") or []),
                    ", ".join(rule.get("rule_id", "") for rule in rules),
                    "Yes" if item.get("requires_review") else "No",
                ]
            )
        _write_sheet(mapping_sheet, headers, rows)

        gap_sheet = workbook.create_sheet("Mapping Gaps")
        _write_sheet(
            gap_sheet,
            ["Gap"],
            [[gap] for gap in (analysis.get("summary_gaps") or [])],
        )

    return _save_workbook(workbook, output_path)
