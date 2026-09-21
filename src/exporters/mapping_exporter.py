"""Export Mapping Agent output as JSON or Excel."""

from __future__ import annotations

import json
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font


def mapping_export_payload(draft: dict[str, Any]) -> dict[str, Any]:
    source_files = draft.get("source_files") or {}
    return {
        "draft_id": draft.get("draft_id", ""),
        "spec_name": source_files.get("spec_name", ""),
        "workflow_mode": source_files.get("workflow_mode", ""),
        "etl_spec": draft.get("etl_spec_draft") or {},
        "mapping_rows": draft.get("mapping_rows") or [],
        "canonical_mappings": draft.get("canonical_mappings") or [],
        "mapping_analysis": draft.get("mapping_analysis") or {},
        "entity_logic": draft.get("entity_logic") or source_files.get("entity_logic") or [],
        "stm_entities": draft.get("stm_entities") or [],
        "document_kind": source_files.get("document_kind") or "",
        "llm_mapping_output": draft.get("llm_mapping_output") or {},
        "missing_information": draft.get("missing_information") or [],
        "assumptions": draft.get("assumptions") or [],
    }


def mapping_json_bytes(draft: dict[str, Any]) -> bytes:
    return json.dumps(mapping_export_payload(draft), indent=2, ensure_ascii=False).encode("utf-8")


def _write_sheet(sheet, headers: list[str], rows: list[list[Any]]) -> None:
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append(row)
    for column_cells in sheet.columns:
        width = min(max(len(str(cell.value or "")) for cell in column_cells) + 2, 80)
        sheet.column_dimensions[column_cells[0].column_letter].width = width


def _analysis_rows(analysis: dict[str, Any]) -> list[list[Any]]:
    rows: list[list[Any]] = []
    for item in analysis.get("mappings") or []:
        source = item.get("source") or {}
        target = item.get("target") or {}
        transform = item.get("transformation") or {}
        rules = item.get("validation_rules") or []
        rows.append(
            [
                item.get("mapping_id", ""),
                source.get("database", ""),
                source.get("schema") or source.get("schema_name", ""),
                source.get("table", ""),
                source.get("column", ""),
                source.get("data_type", ""),
                target.get("database", ""),
                target.get("schema") or target.get("schema_name", ""),
                target.get("table", ""),
                target.get("column", ""),
                target.get("data_type", ""),
                transform.get("type", ""),
                transform.get("logic", ""),
                item.get("risk", ""),
                "; ".join(item.get("gaps") or []),
                ", ".join(rule.get("rule_id", "") for rule in rules),
                "Yes" if item.get("requires_review") else "No",
                item.get("notes", ""),
            ]
        )
    return rows


def mapping_excel_bytes(draft: dict[str, Any]) -> bytes:
    payload = mapping_export_payload(draft)
    workbook = Workbook()

    spec = payload.get("etl_spec") or {}
    summary = workbook.active
    summary.title = "Summary"
    _write_sheet(
        summary,
        ["Field", "Value"],
        [
            ["Draft ID", payload.get("draft_id", "")],
            ["Spec name", payload.get("spec_name", "")],
            ["Workflow mode", payload.get("workflow_mode", "")],
            ["Source table", spec.get("source_table", "")],
            ["Target table", spec.get("target_table", "")],
            ["Load type", spec.get("load_type", "")],
            ["Primary key", spec.get("primary_key", "")],
            ["Incremental column", spec.get("incremental_column", "")],
            ["Mapping rows", len(payload.get("mapping_rows") or [])],
            ["Canonical mappings", len(payload.get("canonical_mappings") or [])],
            ["High risk", (payload.get("mapping_analysis") or {}).get("high_risk_count", 0)],
            ["Gaps", (payload.get("mapping_analysis") or {}).get("gap_count", 0)],
        ],
    )

    row_headers = [
        "Source Database",
        "Source Schema",
        "Source Table",
        "Source Column",
        "Source Data Type",
        "Transformation",
        "Target Database",
        "Target Schema",
        "Target Table",
        "Target Column",
        "Target Data Type",
        "Business Rule",
        "Nullable",
        "Primary Key",
        "Original Source",
    ]
    mapping_sheet = workbook.create_sheet("Mapping Rows")
    _write_sheet(
        mapping_sheet,
        row_headers,
        [
            [
                row.get("source_database", ""),
                row.get("source_schema", ""),
                row.get("source_table", ""),
                row.get("source_column", ""),
                row.get("source_data_type", ""),
                row.get("transformation", ""),
                row.get("target_database", ""),
                row.get("target_schema", ""),
                row.get("target_table", ""),
                row.get("target_column", ""),
                row.get("target_data_type", ""),
                row.get("business_rule", ""),
                row.get("nullable", ""),
                row.get("primary_key", ""),
                row.get("notes", ""),
            ]
            for row in payload.get("mapping_rows") or []
        ],
    )

    canonical_sheet = workbook.create_sheet("Canonical Mappings")
    _write_sheet(
        canonical_sheet,
        [
            "Mapping ID",
            "Source Table",
            "Source Column",
            "Target Table",
            "Target Column",
            "Type",
            "Logic",
            "Business Rule",
            "Notes",
        ],
        [
            [
                item.get("mapping_id", ""),
                (item.get("source") or {}).get("table", ""),
                (item.get("source") or {}).get("column", ""),
                (item.get("target") or {}).get("table", ""),
                (item.get("target") or {}).get("column", ""),
                (item.get("transformation") or {}).get("type", ""),
                (item.get("transformation") or {}).get("logic", ""),
                item.get("business_rule", ""),
                item.get("notes", ""),
            ]
            for item in payload.get("canonical_mappings") or []
        ],
    )

    analysis_sheet = workbook.create_sheet("Mapping Analysis")
    _write_sheet(
        analysis_sheet,
        [
            "Mapping ID",
            "Source Database",
            "Source Schema",
            "Source Table",
            "Source Column",
            "Source Data Type",
            "Target Database",
            "Target Schema",
            "Target Table",
            "Target Column",
            "Target Data Type",
            "Type",
            "Logic",
            "Risk",
            "Gaps",
            "Rule IDs",
            "Requires Review",
            "Notes",
        ],
        _analysis_rows(payload.get("mapping_analysis") or {}),
    )

    gaps_sheet = workbook.create_sheet("Gaps")
    _write_sheet(
        gaps_sheet,
        ["Gap"],
        [[gap] for gap in ((payload.get("mapping_analysis") or {}).get("summary_gaps") or payload.get("missing_information") or [])],
    )

    spec_map = workbook.create_sheet("ETL Spec Mapping")
    _write_sheet(
        spec_map,
        ["Source Column", "Target Column", "Transformation"],
        [
            [row.get("source_column", ""), row.get("target_column", ""), row.get("transformation", "")]
            for row in spec.get("source_to_target_mapping") or []
        ],
    )

    logic_sheet = workbook.create_sheet("Pseudo Code")
    _write_sheet(
        logic_sheet,
        ["Target Table", "Pseudo Code"],
        [
            [item.get("target_table", ""), item.get("pseudo_code", "")]
            for item in payload.get("entity_logic") or []
        ],
    )

    entities = payload.get("stm_entities") or []
    if entities:
        entity_sheet = workbook.create_sheet("STM Entities")
        _write_sheet(
            entity_sheet,
            ["Target Table", "Row Count"],
            [[item.get("target_table", ""), item.get("row_count", "")] for item in entities],
        )

    llm_output = payload.get("llm_mapping_output") or {}
    if llm_output:
        llm_sheet = workbook.create_sheet("LLM Mapping Output")
        llm_rows = [[f"{item.get('target_table', '')} primary key", item.get("primary_key", "")] for item in llm_output.get("entities") or []]
        llm_rows.extend(["Transformation rule", rule] for rule in llm_output.get("transformation_rules") or [])
        llm_rows.extend(["Business rule", rule] for rule in llm_output.get("business_rules") or [])
        _write_sheet(llm_sheet, ["Field", "Value"], llm_rows or [["(none)", ""]])

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()