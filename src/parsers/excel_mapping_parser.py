from pathlib import Path
from typing import Any
import re

from openpyxl import load_workbook

from src.parsers.tabular import coerce_path, is_excel_workbook, is_word_document, read_csv_rows, should_read_as_csv

HEADER_ALIASES = {
    "source_database": {"source database", "src database", "source_db", "source db"},
    "source_schema": {"source schema", "src schema", "source_schema"},
    "source_table": {"source table", "src table", "source_table", "source tbl", "source table(s)", "source tables"},
    "source_column": {
        "source column",
        "src column",
        "source col",
        "source_column",
        "src col",
        "source column & logic",
        "source column and logic",
        "source column / logic",
        "source column logic",
        "source column and logic",
        "src column logic",
        "source field",
        "source field logic",
    },
    "source_data_type": {
        "source data type",
        "source datatype",
        "src data type",
        "source_data_type",
        "source type",
    },
    "target_database": {"target database", "tgt database", "target_db", "target db"},
    "target_schema": {"target schema", "tgt schema", "target_schema"},
    "target_table": {"target table", "tgt table", "target_table", "target tbl"},
    "target_column": {"target column", "tgt column", "target col", "target_column", "tgt col"},
    "target_data_type": {
        "target data type",
        "target datatype",
        "tgt data type",
        "target_data_type",
        "target type",
    },
    "transformation": {"transformation", "transform", "transformation rule", "logic", "mapping logic"},
    "business_rule": {"business rule", "business_rule", "business rules", "rule"},
    "nullable": {"nullable", "null", "nullability", "is nullable"},
    "primary_key": {"primary key", "primary_key", "pk", "is pk"},
    "notes": {"notes", "comment", "comments", "remarks", "original source", "source system"},
}


def _normalize_header(value: Any) -> str:
    text = str(value or "").strip().lower().replace("-", " ").replace("_", " ")
    text = text.replace("&", "and")
    return re.sub(r"\s+", " ", text)


def _alias_lookup() -> dict[str, str]:
    lookup: dict[str, str] = {}
    for canonical, aliases in HEADER_ALIASES.items():
        for alias in aliases:
            lookup[_normalize_header(alias)] = canonical
        lookup[_normalize_header(canonical)] = canonical
    return lookup


def _resolve_columns(headers: list[str]) -> dict[str, int]:
    lookup = _alias_lookup()
    resolved: dict[str, int] = {}
    for index, header in enumerate(headers):
        canonical = lookup.get(_normalize_header(header))
        if canonical and canonical not in resolved:
            resolved[canonical] = index
    required = ("source_column", "target_column")
    missing = [field for field in required if field not in resolved]
    if missing:
        raise ValueError(
            "Mapping sheet must include headers for Source Column and Target Column. "
            f"Found headers: {headers}"
        )
    return resolved


def _qualify(database: str, schema: str, table: str) -> str:
    if "." in (table or ""):
        return table
    parts = [part for part in (database, schema, table) if part]
    return ".".join(parts)


def _cell(row: list[str], column_map: dict[str, int], field: str) -> str:
    if field not in column_map:
        return ""
    index = column_map[field]
    if index >= len(row):
        return ""
    return row[index]


def _parse_mapping_rows(rows: list[Any], path: Path, sheet_name: str) -> dict[str, Any]:
    if not rows:
        raise ValueError(f"Mapping file is empty: {path}")

    headers = [str(cell or "").strip() for cell in rows[0]]
    column_map = _resolve_columns(headers)

    mapping_rows: list[dict[str, str]] = []
    source_table = ""
    target_table = ""

    for raw_row in rows[1:]:
        if not raw_row or all(cell in (None, "") for cell in raw_row):
            continue
        row = ["" if cell is None else str(cell).strip() for cell in raw_row]
        source_column = _cell(row, column_map, "source_column")
        target_column = _cell(row, column_map, "target_column")
        if not source_column and not target_column:
            continue

        row_source_table = _qualify(
            _cell(row, column_map, "source_database"),
            _cell(row, column_map, "source_schema"),
            _cell(row, column_map, "source_table"),
        )
        row_target_table = _qualify(
            _cell(row, column_map, "target_database"),
            _cell(row, column_map, "target_schema"),
            _cell(row, column_map, "target_table"),
        )
        if row_source_table:
            source_table = row_source_table
        if row_target_table:
            target_table = row_target_table

        mapping_rows.append(
            {
                "source_database": _cell(row, column_map, "source_database"),
                "source_schema": _cell(row, column_map, "source_schema"),
                "source_table": row_source_table,
                "source_column": source_column,
                "source_data_type": _cell(row, column_map, "source_data_type"),
                "target_database": _cell(row, column_map, "target_database"),
                "target_schema": _cell(row, column_map, "target_schema"),
                "target_table": row_target_table,
                "target_column": target_column,
                "target_data_type": _cell(row, column_map, "target_data_type"),
                "transformation": _cell(row, column_map, "transformation"),
                "business_rule": _cell(row, column_map, "business_rule"),
                "nullable": _cell(row, column_map, "nullable"),
                "primary_key": _cell(row, column_map, "primary_key"),
                "notes": _cell(row, column_map, "notes"),
            }
        )

    if not mapping_rows:
        raise ValueError(f"No mapping rows found in file: {path}")

    return {
        "source_table": source_table,
        "target_table": target_table,
        "mapping_rows": mapping_rows,
        "sheet_name": sheet_name,
    }


def parse_mapping_workbook(path: Path | str, *, sheet_name: str | None = None) -> dict[str, Any]:
    path = coerce_path(path)
    if is_word_document(path):
        from src.parsers.stm_extract import parse_stm_docx

        return parse_stm_docx(path, write_excel=True)
    if should_read_as_csv(path):
        return _parse_mapping_rows(read_csv_rows(path), path, path.stem)
    if not is_excel_workbook(path):
        raise ValueError(
            f"Unsupported mapping format '{path.name}'. "
            "Use Excel (.xlsx/.xlsm), CSV, or a Word STM document (.docx)."
        )

    workbook = load_workbook(path, data_only=True)
    sheet = workbook[sheet_name] if sheet_name else workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    return _parse_mapping_rows(rows, path, sheet.title)
