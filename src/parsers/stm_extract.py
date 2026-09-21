"""Extract STM tables from Word mapping documents and emit a canonical Excel file.

Word STM headers vary ("Source Column & Logic", "Source Table(s)", ...).
Tables are read with python-docx. Headers are mapped by tokens, not exact labels.
An LLM is used only if a table still has no source/target column after that.
The whole document is never sent to the model.
"""

from __future__ import annotations

import re
import unicodedata
from pathlib import Path
from typing import Any

from src.parsers.tabular import coerce_path

STM_CHAPTER = re.compile(r"source\s+to\s+target\s+mapping", re.I)
PSEUDO_HEADING = re.compile(r"^pseudo[\s\-]*code:?$", re.I)
STM_HEADING = re.compile(r"^stm:?$", re.I)
SKIP_ENTITY = re.compile(r"^(source to target mapping|stm|pseudo code)\b", re.I)
ARCHITECTURE_H2 = re.compile(
    r"^(environments?|lakehouses?|tech stack|orchestrat|"
    r"data ingestion|data processing|data deletion|error handling|"
    r"pipeline load|connections?|data model|document information|"
    r"document purpose|document id)\b",
    re.I,
)
_SQL_START = re.compile(
    r"^(concat|concat_ws|cast|case|coalesce|upper|lower|trim|nvl|ifnull|"
    r"hardcoded|auto|incremental|lookup|n/?a|system|generated)\b",
    re.I,
)

EXCEL_HEADERS = [
    "Source Database",
    "Source Schema",
    "Source Table",
    "Source Column",
    "Source Data Type",
    "Transformation",
    "Target Database",
    "Target Schema",
    "Target Table",
    "Target Column",
    "Target Data Type",
    "Business Rule",
    "Nullable",
    "Primary Key",
    "Notes",
]


def header_tokens(header: str) -> set[str]:
    text = unicodedata.normalize("NFKC", str(header or ""))
    text = text.replace("\xa0", " ").lower()
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return {token for token in text.split() if token}


def canonical_stm_header(header: str) -> str | None:
    tokens = header_tokens(header)
    collapsed = " ".join(token for token in unicodedata.normalize("NFKC", str(header or "")).replace("\xa0", " ").lower().split())
    collapsed = re.sub(r"[^a-z0-9]+", " ", collapsed).strip()
    if "source column" in collapsed:
        return "source_column"
    if "target column" in collapsed:
        return "target_column"
    if "target" in tokens and "column" in tokens:
        return "target_column"
    if "source" in tokens and "column" in tokens:
        return "source_column"
    if "source" in tokens and "table" in tokens:
        return "source_table"
    if "target" in tokens and "table" in tokens:
        return "target_table"
    if "original" in tokens and "source" in tokens:
        return "notes"
    if tokens & {"notes", "comment", "comments", "remarks"}:
        return "notes"
    if "transformation" in tokens or tokens in ({"logic"}, {"mapping", "logic"}):
        return "transformation"
    if "business" in tokens and "rule" in tokens:
        return "business_rule"
    if "source" in tokens and "type" in tokens:
        return "source_data_type"
    if "target" in tokens and "type" in tokens:
        return "target_data_type"
    return None


def resolve_stm_headers(
    headers: list[str],
    *,
    sample_rows: list[list[str]] | None = None,
    use_llm: bool = False,
) -> dict[str, int]:
    resolved: dict[str, int] = {}
    for index, header in enumerate(headers):
        canonical = canonical_stm_header(header)
        if canonical and canonical not in resolved:
            resolved[canonical] = index
    if use_llm and ("source_column" not in resolved or "target_column" not in resolved):
        llm_map = _llm_header_map(headers, sample_rows or [])
        for index, header in enumerate(headers):
            canonical = llm_map.get(header) or llm_map.get(str(header))
            if (
                canonical
                in {
                    "source_column",
                    "target_column",
                    "source_table",
                    "target_table",
                    "transformation",
                    "notes",
                    "source_data_type",
                    "target_data_type",
                    "business_rule",
                }
                and canonical not in resolved
            ):
                resolved[canonical] = index
    return resolved


def is_stm_table(headers: list[str], *, sample_rows: list[list[str]] | None = None) -> bool:
    resolved = resolve_stm_headers(headers, sample_rows=sample_rows, use_llm=False)
    if "target_column" in resolved and "source_column" in resolved:
        return True
    if "target_column" not in resolved:
        return False
    resolved = resolve_stm_headers(headers, sample_rows=sample_rows, use_llm=True)
    return "target_column" in resolved and "source_column" in resolved


def _llm_header_map(headers: list[str], sample_rows: list[list[str]]) -> dict[str, str]:
    tokens: set[str] = set()
    for header in headers:
        tokens |= header_tokens(header)
    if "column" not in tokens or ("source" not in tokens and "target" not in tokens):
        return {}
    try:
        from src.utils.llm_json import invoke_llm_json
    except Exception:
        return {}
    prompt = (
        "Map STM table headers to canonical field names.\n"
        "Canonical fields: source_column, source_table, target_column, target_table, "
        "transformation, notes, source_data_type, target_data_type, business_rule.\n"
        "A header like 'Source Column & Logic' is source_column.\n"
        f"Headers: {headers}\n"
        f"Sample rows: {sample_rows[:3]}\n"
        'Return JSON: {"header_map": {"<exact header>": "<canonical or empty>"}}'
    )
    try:
        payload = invoke_llm_json(
            prompt,
            system_prompt="You map Word STM table headers to canonical ETL mapping columns. Do not invent rows.",
            max_retries=1,
        )
    except Exception:
        return {}
    mapped = payload.get("header_map") or {}
    return {str(key): str(value) for key, value in mapped.items() if value}


def _iter_docx_blocks(path: Path):
    from docx import Document
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(str(path))
    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, document)
        elif isinstance(child, CT_Tbl):
            yield Table(child, document)


def _cell_text(cell) -> str:
    return " ".join((cell.text or "").split()).strip()


def _table_matrix(table) -> list[list[str]]:
    return [[_cell_text(cell) for cell in row.cells] for row in table.rows]


def _heading_level(paragraph) -> int:
    name = (paragraph.style.name if paragraph.style else "") or ""
    match = re.search(r"heading\s*(\d)", name, re.I)
    return int(match.group(1)) if match else 0


def _entity_to_table(name: str) -> str:
    text = (name or "").strip().rstrip(":").strip()
    text = re.sub(r"[^\w\s\-]", "", text)
    return re.sub(r"[\s\-]+", "_", text).strip("_")


def _split_source_and_logic(value: str) -> tuple[str, str]:
    text = (value or "").strip()
    if not text:
        return "", ""
    if _SQL_START.match(text):
        return "", text
    paren = re.match(r"^([A-Za-z_][A-Za-z0-9_]*)\s*\((.+)$", text)
    if paren:
        name = paren.group(1)
        if name.isupper() or name.lower() in {"concat_ws", "concat", "cast", "coalesce"}:
            return "", text
        return name, paren.group(2).strip("() ")
    labeled = re.match(r"^([A-Za-z_][A-Za-z0-9_]*)\s*[:\-–]\s*(.+)$", text)
    if labeled and labeled.group(1).islower():
        return labeled.group(1), labeled.group(2).strip()
    if re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", text):
        return text, "DIRECT"
    return text, text


def write_normalized_mapping_excel(
    mapping_rows: list[dict[str, str]],
    *,
    output_path: Path,
    entity_logic: list[dict[str, str]] | None = None,
) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Mapping"
    sheet.append(EXCEL_HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in mapping_rows:
        sheet.append(
            [
                row.get("source_database", ""),
                row.get("source_schema", ""),
                row.get("source_table", ""),
                row.get("source_column", ""),
                row.get("source_data_type", ""),
                row.get("transformation", ""),
                row.get("target_database", ""),
                row.get("target_schema", ""),
                row.get("target_table", ""),
                row.get("target_column", ""),
                row.get("target_data_type", ""),
                row.get("business_rule", ""),
                row.get("nullable", ""),
                row.get("primary_key", ""),
                row.get("notes", ""),
            ]
        )
    if entity_logic:
        logic_sheet = workbook.create_sheet("Entity Logic")
        logic_sheet.append(["Target Table", "Pseudo Code"])
        for item in entity_logic:
            logic_sheet.append([item.get("target_table", ""), item.get("pseudo_code", "")])
    workbook.save(output_path)
    return output_path


def stm_entity_summary(mapping_rows: list[dict[str, str]]) -> list[dict[str, Any]]:
    counts: dict[str, int] = {}
    for row in mapping_rows:
        name = str(row.get("target_table") or "").strip() or "(unscoped)"
        counts[name] = counts.get(name, 0) + 1
    return [{"target_table": name, "row_count": count} for name, count in counts.items()]


def parse_stm_docx(path: Path | str, *, write_excel: bool = False) -> dict[str, Any]:
    path = coerce_path(path)
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    current_entity = ""
    current_section = ""
    in_stm_chapter = False
    saw_stm_chapter = False
    entity_logic: dict[str, list[str]] = {}
    mapping_rows: list[dict[str, str]] = []

    def table_name() -> str:
        return _entity_to_table(current_entity)

    def accept_stm_content() -> bool:
        return in_stm_chapter or not saw_stm_chapter

    for block in _iter_docx_blocks(path):
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if not text:
                continue
            level = _heading_level(block)
            if level == 1:
                if STM_CHAPTER.search(text):
                    in_stm_chapter = True
                    saw_stm_chapter = True
                    current_entity = ""
                    current_section = ""
                else:
                    in_stm_chapter = False
                    current_entity = ""
                    current_section = ""
                continue
            if level == 2:
                if not accept_stm_content():
                    continue
                if SKIP_ENTITY.match(text) or PSEUDO_HEADING.match(text) or STM_HEADING.match(text):
                    continue
                if ARCHITECTURE_H2.match(text):
                    current_entity = ""
                    current_section = ""
                    continue
                current_entity = text
                current_section = ""
                entity_logic.setdefault(table_name(), [])
                continue
            if level == 3 or STM_HEADING.match(text) or PSEUDO_HEADING.match(text):
                if not accept_stm_content():
                    continue
                if PSEUDO_HEADING.match(text):
                    current_section = "pseudo"
                elif STM_HEADING.match(text):
                    current_section = "stm"
                continue
            if current_section == "pseudo" and current_entity and accept_stm_content():
                entity_logic.setdefault(table_name(), []).append(text)
            continue

        if not isinstance(block, Table):
            continue
        if not accept_stm_content() or not current_entity:
            continue

        matrix = _table_matrix(block)
        if len(matrix) < 2:
            continue
        headers = matrix[0]
        body = matrix[1:]
        if not is_stm_table(headers, sample_rows=body[:3]):
            continue

        column_map = resolve_stm_headers(headers, sample_rows=body[:3], use_llm=True)
        target_table = table_name()
        for raw in body:
            if not raw or all(not cell for cell in raw):
                continue
            target_column = raw[column_map["target_column"]] if column_map["target_column"] < len(raw) else ""
            source_raw = raw[column_map["source_column"]] if column_map["source_column"] < len(raw) else ""
            source_table = ""
            if "source_table" in column_map and column_map["source_table"] < len(raw):
                source_table = raw[column_map["source_table"]]
            notes = ""
            if "notes" in column_map and column_map["notes"] < len(raw):
                notes = raw[column_map["notes"]]
            if not target_column and not source_raw:
                continue
            source_column, transformation = _split_source_and_logic(source_raw)
            mapping_rows.append(
                {
                    "source_database": "",
                    "source_schema": "",
                    "source_table": source_table,
                    "source_column": source_column or source_raw,
                    "source_data_type": "",
                    "target_database": "",
                    "target_schema": "",
                    "target_table": target_table,
                    "target_column": target_column,
                    "target_data_type": "",
                    "transformation": transformation,
                    "business_rule": "",
                    "nullable": "",
                    "primary_key": "",
                    "notes": notes,
                }
            )

    if not mapping_rows:
        raise ValueError(
            f"No source-to-target mapping tables found in Word document: {path}. "
            "Expected STM tables with a target column and a source column (or Source Column & Logic)."
        )

    logic_payload = [
        {"target_table": name, "pseudo_code": "\n".join(lines)}
        for name, lines in entity_logic.items()
        if name and lines
    ]
    entities = stm_entity_summary(mapping_rows)
    first_source = next((row["source_table"] for row in mapping_rows if row["source_table"]), "")
    first_target = next((row["target_table"] for row in mapping_rows if row["target_table"]), "")
    payload: dict[str, Any] = {
        "source_table": first_source,
        "target_table": first_target,
        "mapping_rows": mapping_rows,
        "sheet_name": "STM",
        "entity_logic": logic_payload,
        "stm_entities": entities,
        "document_kind": "word_stm",
        "normalized_excel_path": "",
    }
    if write_excel:
        excel_path = path.with_name(f"{path.stem}_normalized_mapping.xlsx")
        write_normalized_mapping_excel(mapping_rows, output_path=excel_path, entity_logic=logic_payload)
        payload["normalized_excel_path"] = str(excel_path)
    return payload
