from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from src.parsers.document_parser import read_requirement_document
from src.parsers.tabular import coerce_path, read_csv_rows, should_read_as_csv

KPI_HEADER_ALIASES = {
    "kpi_id": {"kpi id", "id", "kpi_id", "metric id", "measure id"},
    "kpi_name": {"kpi name", "kpi", "metric", "measure", "kpi title", "name"},
    "module": {"module", "screen", "page", "feature", "ui screen"},
    "sql_query": {
        "sql query",
        "query",
        "sql",
        "logic",
        "databricks sql",
        "validation sql",
        "kpi sql",
        "logic query",
    },
    "expected_result": {
        "expected result",
        "expected",
        "pass criteria",
        "expected output",
        "validation result",
    },
    "priority": {"priority", "severity"},
}


def _normalize_header(value: Any) -> str:
    return str(value or "").strip().lower()


def _resolve_kpi_columns(headers: list[str]) -> dict[str, int]:
    resolved: dict[str, int] = {}
    for index, header in enumerate(headers):
        normalized = _normalize_header(header)
        for canonical, aliases in KPI_HEADER_ALIASES.items():
            if normalized in aliases and canonical not in resolved:
                resolved[canonical] = index
    if "sql_query" not in resolved:
        raise ValueError(
            "KPI logic sheet must include a SQL/Query column. "
            f"Found headers: {headers}"
        )
    return resolved


def _clean_sql(value: str) -> str:
    text = (value or "").strip()
    if text.startswith("```"):
        text = re.sub(r"^```(?:sql)?\s*", "", text, flags=re.IGNORECASE)
        text = re.sub(r"\s*```$", "", text)
    return text.strip()


def _entry(
    *,
    kpi_id: str,
    kpi_name: str,
    sql_query: str,
    expected_result: str = "",
    module: str = "",
    priority: str = "High",
    source_row: int | None = None,
) -> dict[str, Any]:
    return {
        "kpi_id": kpi_id or "UNASSIGNED",
        "kpi_name": kpi_name or kpi_id or "Unnamed KPI",
        "module": module,
        "sql_query": _clean_sql(sql_query),
        "expected_result": expected_result or "0 rows returned (no validation failures).",
        "priority": priority or "High",
        "source_row": source_row,
    }


def _parse_kpi_logic_rows(rows: list[Any], path: Path) -> list[dict[str, Any]]:
    if not rows:
        raise ValueError(f"KPI logic file is empty: {path}")

    headers = [str(cell or "").strip() for cell in rows[0]]
    column_map = _resolve_kpi_columns(headers)
    entries: list[dict[str, Any]] = []

    for row_number, raw_row in enumerate(rows[1:], start=2):
        if not raw_row or all(cell in (None, "") for cell in raw_row):
            continue
        row = ["" if cell is None else str(cell).strip() for cell in raw_row]
        if column_map["sql_query"] >= len(row):
            continue
        sql_query = row[column_map["sql_query"]]
        if not sql_query:
            continue

        def _col(name: str) -> str:
            index = column_map.get(name)
            if index is None or index >= len(row):
                return ""
            return row[index]

        kpi_name = _col("kpi_name")
        kpi_id = _col("kpi_id") or f"KPI-{len(entries) + 1:03d}"
        if not kpi_name:
            kpi_name = kpi_id

        entries.append(
            _entry(
                kpi_id=kpi_id,
                kpi_name=kpi_name,
                sql_query=sql_query,
                expected_result=_col("expected_result"),
                module=_col("module"),
                priority=_col("priority") or "High",
                source_row=row_number,
            )
        )

    if not entries:
        raise ValueError(f"No KPI SQL rows found in file: {path}")
    return entries


def parse_kpi_logic_excel(path: Path, *, sheet_name: str | None = None) -> list[dict[str, Any]]:
    path = coerce_path(path)
    if should_read_as_csv(path):
        return _parse_kpi_logic_rows(read_csv_rows(path), path)
    try:
        workbook = load_workbook(path, data_only=True)
    except Exception:
        return _parse_kpi_logic_rows(read_csv_rows(path), path)
    sheet = workbook[sheet_name] if sheet_name else workbook.active
    return _parse_kpi_logic_rows(list(sheet.iter_rows(values_only=True)), path)


def parse_kpi_logic_sql(path: Path) -> list[dict[str, Any]]:
    text = path.read_text(encoding="utf-8")
    entries: list[dict[str, Any]] = []
    blocks = re.split(r"(?m)^\s*--\s*KPI\s*:\s*", text)
    if len(blocks) > 1:
        for block in blocks[1:]:
            lines = block.strip().splitlines()
            if not lines:
                continue
            kpi_name = lines[0].strip()
            sql_query = "\n".join(lines[1:]).strip()
            if sql_query:
                entries.append(
                    _entry(
                        kpi_id=f"KPI-{len(entries) + 1:03d}",
                        kpi_name=kpi_name,
                        sql_query=sql_query,
                    )
                )
        if entries:
            return entries

    select_blocks = re.findall(
        r"(?:/\*.*?\*/\s*)?(SELECT[\s\S]*?)(?=\n\s*(?:--\s*KPI|SELECT|/\*)|\Z)",
        text,
        flags=re.IGNORECASE,
    )
    for index, sql_query in enumerate(select_blocks, start=1):
        entries.append(
            _entry(
                kpi_id=f"KPI-{index:03d}",
                kpi_name=f"KPI Query {index}",
                sql_query=sql_query,
            )
        )

    if not entries:
        raise ValueError(f"No SELECT queries found in SQL file: {path}")
    return entries


def parse_kpi_logic_text(path: Path) -> list[dict[str, Any]]:
    text = read_requirement_document(path) if path.suffix.lower() != ".txt" else path.read_text(encoding="utf-8")
    entries: list[dict[str, Any]] = []

    pattern = re.compile(
        r"(?is)(?:^|\n)\s*(?:KPI\s*[:#-]\s*)?(?P<name>[A-Za-z0-9 \-_/]+?)\s*(?:\n|$)(?:Expected\s*[:=-].*?\n)?(?P<sql>(?:WITH|SELECT)[\s\S]*?)(?=\n\s*(?:KPI\s*[:#-]|$))",
        re.MULTILINE,
    )
    for index, match in enumerate(pattern.finditer(text), start=1):
        entries.append(
            _entry(
                kpi_id=f"KPI-{index:03d}",
                kpi_name=match.group("name").strip(),
                sql_query=match.group("sql").strip(),
            )
        )

    if entries:
        return entries

    select_blocks = re.findall(r"((?:WITH|SELECT)[\s\S]*?)(?=\n\s*(?:WITH|SELECT)\b|\Z)", text, flags=re.IGNORECASE)
    for index, sql_query in enumerate(select_blocks, start=1):
        entries.append(
            _entry(
                kpi_id=f"KPI-{index:03d}",
                kpi_name=f"KPI Query {index}",
                sql_query=sql_query.strip(),
            )
        )

    if not entries:
        raise ValueError(f"No KPI SQL blocks found in document: {path}")
    return entries


def parse_kpi_logic_document(path: Path) -> list[dict[str, Any]]:
    path = coerce_path(path)
    suffix = path.suffix.lower()
    if should_read_as_csv(path) or suffix == ".csv":
        return _parse_kpi_logic_rows(read_csv_rows(path), path)
    if suffix in {".xlsx", ".xlsm"}:
        return parse_kpi_logic_excel(path)
    if suffix == ".sql":
        return parse_kpi_logic_sql(path)
    if suffix in {".txt", ".docx", ".pdf", ".md"}:
        return parse_kpi_logic_text(path)
    if suffix == ".json":
        payload = json.loads(path.read_text(encoding="utf-8"))
        if isinstance(payload, list):
            return payload
        if isinstance(payload, dict) and isinstance(payload.get("queries"), list):
            return payload["queries"]
        raise ValueError("KPI logic JSON must be an array or an object with a queries array.")
    raise ValueError(
        f"Unsupported KPI logic document format: {suffix}. "
        "Use .xlsx, .csv, .sql, .txt, .docx, .pdf, .md, or .json."
    )
